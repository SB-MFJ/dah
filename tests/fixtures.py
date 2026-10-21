"""Generate small test documents in-memory."""

from __future__ import annotations

import io


def make_pdf(pages: list[str]) -> bytes:
    """Minimal valid PDF with one Helvetica text line per page (empty string => no text layer)."""
    objs: list[bytes] = []
    n = len(pages)
    kids = " ".join(f"{3 + i * 2} 0 R" for i in range(n))
    objs.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objs.append(f"<< /Type /Pages /Kids [{kids}] /Count {n} >>".encode())
    font_id = 3 + n * 2
    for i, text in enumerate(pages):
        content_id = 4 + i * 2
        objs.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {content_id} 0 R "
            f"/Resources << /Font << /F1 {font_id} 0 R >> >> >>".encode()
        )
        safe = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        stream = f"BT /F1 18 Tf 72 700 Td ({safe}) Tj ET".encode() if text else b""
        objs.append(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
    objs.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for i, o in enumerate(objs, start=1):
        offsets.append(out.tell())
        out.write(f"{i} 0 obj\n".encode() + o + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode())
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return out.getvalue()


def make_docx(paragraphs: list[str], table: list[list[str]] | None = None, author: str = "Ada") -> bytes:
    import docx

    d = docx.Document()
    d.core_properties.author = author
    d.core_properties.title = paragraphs[0] if paragraphs else ""
    for p in paragraphs:
        d.add_paragraph(p)
    if table:
        t = d.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, val in enumerate(row):
                t.cell(r, c).text = val
    b = io.BytesIO()
    d.save(b)
    return b.getvalue()


def make_xlsx(rows: list[list]) -> bytes:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales"
    for r in rows:
        ws.append(r)
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def make_png(text: str = "HELLO OCR") -> bytes:
    from PIL import Image, ImageDraw, ImageFont

    im = Image.new("RGB", (900, 200), "white")
    dr = ImageDraw.Draw(im)
    try:
        font = ImageFont.load_default(size=72)
    except TypeError:  # pragma: no cover - old Pillow
        font = ImageFont.load_default()
    dr.text((30, 50), text, fill="black", font=font)
    b = io.BytesIO()
    im.save(b, format="PNG")
    return b.getvalue()
