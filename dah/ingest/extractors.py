"""Text + metadata extraction per file type, with OCR for images and scanned PDF pages.

Each extractor returns an ``Extraction``: a list of pages ``(text, was_ocr)``, flat metadata
(common keys: title, author, subject, keywords, created, modified) plus format-specific keys,
and optional tabular ``records`` ``(sheet, row_number, {column: value})``.
"""

from __future__ import annotations

import csv
import datetime as dt
import email
import email.policy
import io
import json
import logging
import shutil
import subprocess
import tempfile
from dataclasses import dataclass, field
from html.parser import HTMLParser
from pathlib import Path
from typing import Callable

from ..config import get_settings

log = logging.getLogger("dah.ingest")

MAX_RECORDS = 1_000_000


@dataclass
class Extraction:
    pages: list[tuple[str, bool]] = field(default_factory=list)
    metadata: dict = field(default_factory=dict)
    records: list[tuple[str, int, dict]] = field(default_factory=list)
    method: str = "none"
    warnings: list[str] = field(default_factory=list)

    @property
    def text(self) -> str:
        return "\n\n".join(t for t, _ in self.pages if t)


class ExtractorUnavailable(RuntimeError):
    pass


def _iso(v) -> str | None:
    if v is None or v == "":
        return None
    if isinstance(v, (dt.datetime, dt.date)):
        return v.isoformat()
    return str(v)


def _clean(d: dict) -> dict:
    return {k: v for k, v in d.items() if v not in (None, "", [], {})}


# --------------------------------------------------------------------------- OCR


def ocr_available() -> bool:
    if not shutil.which("tesseract"):
        return False
    try:
        import pytesseract  # noqa: F401
    except ImportError:  # pragma: no cover
        return False
    return True


def ocr_image(img) -> tuple[str, float | None]:
    """OCR a PIL image. Returns (text, mean word confidence 0-100)."""
    import pytesseract
    from PIL import ImageOps

    img = ImageOps.exif_transpose(img)
    if img.mode not in ("L", "RGB"):
        img = img.convert("RGB")
    gray = img.convert("L")
    if max(gray.size) < 1200:  # upscale small images; tesseract likes ~300dpi glyphs
        f = 1200 / max(gray.size)
        gray = gray.resize((int(gray.width * f), int(gray.height * f)))
    data = pytesseract.image_to_data(gray, lang=get_settings().ocr_languages, output_type=pytesseract.Output.DICT)
    lines: dict[tuple, list[str]] = {}
    confs = []
    for i, word in enumerate(data["text"]):
        if not word.strip():
            continue
        key = (data["block_num"][i], data["par_num"][i], data["line_num"][i])
        lines.setdefault(key, []).append(word)
        try:
            c = float(data["conf"][i])
            if c >= 0:
                confs.append(c)
        except (TypeError, ValueError):
            pass
    text = "\n".join(" ".join(ws) for _, ws in sorted(lines.items()))
    return text, (round(sum(confs) / len(confs), 1) if confs else None)


# --------------------------------------------------------------------------- text-ish


def _decode(raw: bytes) -> tuple[str, str]:
    from charset_normalizer import from_bytes

    best = from_bytes(raw).best()
    if best is None:
        return raw.decode("utf-8", errors="replace"), "utf-8"
    return str(best), best.encoding


def extract_text(path: Path, mime: str) -> Extraction:
    text, enc = _decode(path.read_bytes())
    pages = [p for p in text.split("\f")] if "\f" in text else [text]
    md: dict = {"encoding": enc, "line_count": text.count("\n") + (1 if text else 0)}
    if mime == "text/markdown":
        for line in text.splitlines():
            if line.startswith("# "):
                md["title"] = line[2:].strip()
                break
    return Extraction(pages=[(p, False) for p in pages], metadata=md, method="text")


def extract_csv(path: Path, mime: str) -> Extraction:
    text, enc = _decode(path.read_bytes())
    sample = text[:65536]
    delim = "\t" if mime == "text/tab-separated-values" else None
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=delim or ",;\t|")
    except csv.Error:
        dialect = csv.excel_tab if delim else csv.excel
    reader = csv.DictReader(io.StringIO(text), dialect=dialect)
    records = []
    for i, row in enumerate(reader, start=1):
        if i > MAX_RECORDS:
            break
        records.append(("csv", i, {(k or f"col{j}"): v for j, (k, v) in enumerate(row.items())}))
    md = {"encoding": enc, "columns": reader.fieldnames or [], "row_count": len(records), "delimiter": dialect.delimiter}
    return Extraction(pages=[(text, False)], metadata=md, records=records, method="csv")


def extract_json(path: Path, mime: str) -> Extraction:
    text, enc = _decode(path.read_bytes())
    records: list = []
    md: dict = {"encoding": enc}
    try:
        if mime == "application/x-ndjson":
            items = [json.loads(line) for line in text.splitlines() if line.strip()]
        else:
            items = json.loads(text)
        if isinstance(items, dict):
            md["top_level_keys"] = sorted(items)[:200]
            items = [items]
        if isinstance(items, list):
            for i, it in enumerate(items[:MAX_RECORDS], start=1):
                records.append(("json", i, it if isinstance(it, dict) else {"value": it}))
            md["row_count"] = len(records)
    except ValueError as e:
        return Extraction(pages=[(text, False)], metadata=md, method="text", warnings=[f"invalid JSON: {e}"])
    return Extraction(pages=[(text, False)], metadata=md, records=records, method="json")


class _HTMLText(HTMLParser):
    SKIP = {"script", "style", "noscript", "template"}
    BLOCK = {"p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6", "section", "article", "table"}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self.skip = 0
        self.title = ""
        self._in_title = False
        self.meta: dict = {}

    def handle_starttag(self, tag, attrs):
        if tag in self.SKIP:
            self.skip += 1
        if tag == "title":
            self._in_title = True
        if tag == "meta":
            a = dict(attrs)
            if a.get("name") and a.get("content"):
                self.meta[a["name"].lower()] = a["content"]
        if tag in self.BLOCK:
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in self.SKIP and self.skip:
            self.skip -= 1
        if tag == "title":
            self._in_title = False

    def handle_data(self, data):
        if self._in_title:
            self.title += data
        elif not self.skip:
            self.parts.append(data)

    def text(self) -> str:
        raw = "".join(self.parts)
        lines = [" ".join(line.split()) for line in raw.splitlines()]
        return "\n".join(line for line in lines if line)


def html_to_text(html: str) -> tuple[str, dict]:
    p = _HTMLText()
    p.feed(html)
    md = {"title": p.title.strip() or None, "author": p.meta.get("author"), "description": p.meta.get("description"), "keywords": p.meta.get("keywords")}
    return p.text(), _clean(md)


def extract_html(path: Path, mime: str) -> Extraction:
    raw, enc = _decode(path.read_bytes())
    text, md = html_to_text(raw)
    md["encoding"] = enc
    return Extraction(pages=[(text, False)], metadata=md, method="html")


def extract_email(path: Path, mime: str) -> Extraction:
    msg = email.message_from_bytes(path.read_bytes(), policy=email.policy.default)
    body = msg.get_body(preferencelist=("plain", "html"))
    text = ""
    if body is not None:
        content = body.get_content()
        text = html_to_text(content)[0] if body.get_content_type() == "text/html" else content
    attachments = [a.get_filename() for a in msg.iter_attachments() if a.get_filename()]
    md = _clean(
        {
            "title": msg.get("subject"),
            "subject": msg.get("subject"),
            "author": msg.get("from"),
            "from": msg.get("from"),
            "to": msg.get("to"),
            "cc": msg.get("cc"),
            "created": msg.get("date"),
            "message_id": msg.get("message-id"),
            "attachments": attachments,
        }
    )
    return Extraction(pages=[(text, False)], metadata=md, method="email")


# --------------------------------------------------------------------------- PDF


def extract_pdf(path: Path, mime: str) -> Extraction:
    from pypdf import PdfReader

    settings = get_settings()
    reader = PdfReader(str(path))
    md: dict = {"encrypted": reader.is_encrypted}
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception:
            return Extraction(metadata=md, method="pdf", warnings=["encrypted PDF; no password"])
    info = reader.metadata or {}
    md.update(
        _clean(
            {
                "title": info.get("/Title"),
                "author": info.get("/Author"),
                "subject": info.get("/Subject"),
                "keywords": info.get("/Keywords"),
                "creator": info.get("/Creator"),
                "producer": info.get("/Producer"),
                "created": _iso(getattr(info, "creation_date", None)),
                "modified": _iso(getattr(info, "modification_date", None)),
            }
        )
    )
    md = {k: (str(v) if not isinstance(v, (bool, int, float, list, dict)) else v) for k, v in md.items()}

    can_ocr = ocr_available()
    pages: list[tuple[str, bool]] = []
    warnings: list[str] = []
    confs: list[float] = []
    pdfium_doc = None
    for i, page in enumerate(reader.pages):
        try:
            text = page.extract_text() or ""
        except Exception as e:  # malformed content streams happen
            text = ""
            warnings.append(f"page {i + 1}: text layer error {e}")
        if len(text.strip()) >= settings.ocr_min_page_chars:
            pages.append((text, False))
            continue
        if not can_ocr:
            pages.append((text, False))
            continue
        try:
            import pypdfium2 as pdfium

            if pdfium_doc is None:
                pdfium_doc = pdfium.PdfDocument(str(path))
            img = pdfium_doc[i].render(scale=300 / 72).to_pil()
            otext, conf = ocr_image(img)
            if conf is not None:
                confs.append(conf)
            pages.append((otext if len(otext.strip()) > len(text.strip()) else text, True))
        except Exception as e:
            warnings.append(f"page {i + 1}: OCR failed {e}")
            pages.append((text, False))
    if pdfium_doc is not None:
        pdfium_doc.close()
    ocr_pages = sum(1 for _, o in pages if o)
    if ocr_pages and confs:
        md["ocr_confidence"] = round(sum(confs) / len(confs), 1)
    if not can_ocr and any(len(t.strip()) < settings.ocr_min_page_chars for t, _ in pages):
        warnings.append("some pages have no text layer and OCR (tesseract) is not installed")
    return Extraction(pages=pages, metadata=md, method="pdf+ocr" if ocr_pages else "pdf", warnings=warnings)


# --------------------------------------------------------------------------- Office (OOXML)


def _core_props(cp) -> dict:
    """OOXML core properties (python-docx/python-pptx and openpyxl name a few fields differently)."""
    g = lambda *names: next((getattr(cp, n) for n in names if getattr(cp, n, None)), None)  # noqa: E731
    return _clean(
        {
            "title": g("title"),
            "author": g("author", "creator"),
            "subject": g("subject"),
            "keywords": g("keywords"),
            "category": g("category"),
            "comments": g("comments", "description"),
            "last_modified_by": g("last_modified_by", "lastModifiedBy"),
            "revision": g("revision"),
            "created": _iso(g("created")),
            "modified": _iso(g("modified")),
        }
    )


def extract_docx(path: Path, mime: str) -> Extraction:
    import docx

    d = docx.Document(str(path))
    parts = [p.text for p in d.paragraphs]
    records = []
    for ti, table in enumerate(d.tables, start=1):
        rows = [[c.text.strip() for c in r.cells] for r in table.rows]
        parts.append("\n".join("\t".join(r) for r in rows))
        if len(rows) > 1:
            header = [h or f"col{j}" for j, h in enumerate(rows[0])]
            for ri, r in enumerate(rows[1:], start=1):
                records.append((f"table{ti}", ri, dict(zip(header, r))))
    for section in d.sections:
        for hf in (section.header, section.footer):
            t = "\n".join(p.text for p in hf.paragraphs if p.text.strip())
            if t:
                parts.append(t)
    md = _core_props(d.core_properties)
    md["table_count"] = len(d.tables)
    md["paragraph_count"] = len(d.paragraphs)
    return Extraction(pages=[("\n".join(parts), False)], metadata=md, records=records, method="docx")


def extract_pptx(path: Path, mime: str) -> Extraction:
    from pptx import Presentation

    prs = Presentation(str(path))
    pages = []
    for slide in prs.slides:
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                texts.append(shape.text_frame.text)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    texts.append("\t".join(c.text for c in row.cells))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text
            if notes.strip():
                texts.append("Notes: " + notes)
        pages.append(("\n".join(texts), False))
    md = _core_props(prs.core_properties)
    md["slide_count"] = len(pages)
    return Extraction(pages=pages, metadata=md, method="pptx")


def extract_xlsx(path: Path, mime: str) -> Extraction:
    import openpyxl

    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    pages, records = [], []
    sheet_rows = {}
    for ws in wb.worksheets:
        rows_iter = ws.iter_rows(values_only=True)
        header = None
        lines = []
        n = 0
        for row in rows_iter:
            if all(v is None for v in row):
                continue
            lines.append("\t".join("" if v is None else str(v) for v in row))
            if header is None:
                header = [str(h) if h is not None else f"col{j}" for j, h in enumerate(row)]
                continue
            n += 1
            if len(records) < MAX_RECORDS:
                records.append((ws.title, n, {header[j] if j < len(header) else f"col{j}": _iso(v) if isinstance(v, (dt.date, dt.datetime)) else v for j, v in enumerate(row)}))
        sheet_rows[ws.title] = n
        pages.append((f"# {ws.title}\n" + "\n".join(lines), False))
    md = _core_props(wb.properties)
    md.update({"sheet_names": list(sheet_rows), "row_count": sum(sheet_rows.values())})
    wb.close()
    return Extraction(pages=pages, metadata=md, records=records, method="xlsx")


# --------------------------------------------------------------------------- legacy via LibreOffice

LEGACY_TARGET = {
    "application/msword": ("docx", extract_docx),
    "application/rtf": ("docx", extract_docx),
    "text/rtf": ("docx", extract_docx),
    "application/vnd.oasis.opendocument.text": ("docx", extract_docx),
    "application/vnd.wordperfect": ("docx", extract_docx),
    "application/vnd.ms-excel": ("xlsx", extract_xlsx),
    "application/vnd.oasis.opendocument.spreadsheet": ("xlsx", extract_xlsx),
    "application/vnd.ms-powerpoint": ("pptx", extract_pptx),
    "application/vnd.oasis.opendocument.presentation": ("pptx", extract_pptx),
}


def extract_legacy(path: Path, mime: str) -> Extraction:
    target, fn = LEGACY_TARGET[mime]
    soffice = shutil.which(get_settings().soffice_bin) or shutil.which("libreoffice")
    if soffice:
        with tempfile.TemporaryDirectory(prefix="dah-soffice-") as td:
            # Isolated profile dir so concurrent conversions don't fight over a lock.
            profile = Path(td) / "profile"
            cmd = [soffice, f"-env:UserInstallation=file://{profile}", "--headless", "--convert-to", target, "--outdir", td, str(path)]
            subprocess.run(cmd, check=True, capture_output=True, timeout=300)
            out = next(Path(td).glob(f"*.{target}"), None)
            if out is None:
                raise ExtractorUnavailable(f"LibreOffice produced no .{target}")
            ex = fn(out, mime)
            ex.method = f"libreoffice->{ex.method}"
            return ex
    if mime == "application/msword" and shutil.which("antiword"):
        out = subprocess.run(["antiword", str(path)], check=True, capture_output=True, timeout=120).stdout
        return Extraction(pages=[(out.decode("utf-8", "replace"), False)], method="antiword")
    raise ExtractorUnavailable(f"no converter for {mime}: install LibreOffice (included in the Docker image)")


# --------------------------------------------------------------------------- images


def extract_image(path: Path, mime: str) -> Extraction:
    from PIL import ExifTags, Image, ImageSequence

    md: dict = {}
    warnings: list[str] = []
    pages: list[tuple[str, bool]] = []
    with Image.open(path) as im:
        md.update({"width": im.width, "height": im.height, "image_format": im.format, "color_mode": im.mode})
        if im.info.get("dpi"):
            md["dpi"] = [float(x) for x in im.info["dpi"]]
        frames = getattr(im, "n_frames", 1)
        md["frames"] = frames
        try:
            exif = im.getexif()
            tags = {}
            for k, v in exif.items():
                name = ExifTags.TAGS.get(k, str(k))
                if isinstance(v, bytes):
                    continue
                tags[name] = v if isinstance(v, (int, float, str)) else str(v)
            if tags:
                md["exif"] = tags
                md.update(_clean({"created": tags.get("DateTime"), "author": tags.get("Artist"), "camera": tags.get("Model")}))
        except Exception:  # pragma: no cover
            pass
        if ocr_available():
            confs = []
            for frame in ImageSequence.Iterator(im):
                text, conf = ocr_image(frame.copy())
                pages.append((text, True))
                if conf is not None:
                    confs.append(conf)
                if len(pages) >= 500:
                    warnings.append("stopped OCR after 500 frames")
                    break
            if confs:
                md["ocr_confidence"] = round(sum(confs) / len(confs), 1)
        else:
            warnings.append("OCR (tesseract) is not installed; image text not extracted")
    return Extraction(pages=pages, metadata=md, method="ocr" if pages else "image", warnings=warnings)


# --------------------------------------------------------------------------- dispatch

OOXML_DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
OOXML_XLSX = ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "application/vnd.ms-excel.sheet.macroEnabled.12")
OOXML_PPTX = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


def pick_extractor(mime: str) -> Callable[[Path, str], Extraction] | None:
    if mime == "application/pdf":
        return extract_pdf
    if mime == OOXML_DOCX:
        return extract_docx
    if mime in OOXML_XLSX:
        return extract_xlsx
    if mime == OOXML_PPTX:
        return extract_pptx
    if mime in LEGACY_TARGET:
        return extract_legacy
    if mime.startswith("image/") and mime not in ("image/svg+xml",):
        return extract_image
    if mime in ("text/csv", "text/tab-separated-values"):
        return extract_csv
    if mime in ("application/json", "application/x-ndjson"):
        return extract_json
    if mime in ("text/html", "application/xhtml+xml", "image/svg+xml"):
        return extract_html
    if mime == "message/rfc822":
        return extract_email
    if mime.startswith("text/") or mime in ("application/xml", "application/yaml", "application/x-yaml"):
        return extract_text
    return None


def extract(path: Path, mime: str) -> Extraction:
    fn = pick_extractor(mime)
    if fn is None:
        from .detect import looks_like_text

        with open(path, "rb") as f:
            if looks_like_text(f.read(8192)):
                return extract_text(path, "text/plain")
        return Extraction(method="none", warnings=[f"no extractor for {mime}; stored with metadata only"])
    return fn(path, mime)
